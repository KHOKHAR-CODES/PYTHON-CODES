"""Loader for Word (.docx) files."""

from typing import List, Dict

from docx import Document


def load_docx(file_path: str) -> List[Dict]:
    """Load a .docx file and return its text as a single record.

    Word documents don't have a simple 'page' concept via python-docx, so
    page is left as None; the text splitter will chunk it further.
    """
    doc = Document(file_path)
    paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]

    for table in doc.tables:
        for row in table.rows:
            row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
            if row_text:
                paragraphs.append(row_text)

    text = "\n".join(paragraphs).strip()
    if not text:
        return []
    return [{"text": text, "source": file_path, "page": None}]
